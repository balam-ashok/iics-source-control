"""
IDMC: migrate connections or schedules from the Demo org to the Dfactory org on demand, without a check-in.

Two ways to say what to migrate:
  excel  rows marked Migrate = Y in cicd/migration_objects.xlsx (sheets "Connections" and "Schedules");
         runs automatically when that file is pushed to main, or by hand
  names  a comma-separated list typed into the Run workflow form (* = every one in Demo)

Run by .github/workflows/idmc_migrate_objects.yml. Reuses idmc_cicd.py for the logins, the org check and
the connection logic, and cicd/config.yml for connection_map, runtime_environment_map, connection
passwords/overrides and the expected org names. Writes migration_results.xlsx and a run summary.

Environment:
  IICS_LOGIN_URL, IICS_USERNAME / IICS_PASSWORD (Demo), UAT_IICS_USERNAME / UAT_IICS_PASSWORD (Dfactory)
  SOURCE            excel | names (default: names when NAMES is set, else excel)
  OBJECT_TYPE       connection | schedule | both (both = excel only)
  NAMES             names mode: names as in Demo, comma or newline separated; * = all
  UPDATE_EXISTING   names mode: overwrite Dfactory copies whose definition differs (default false)
  TEST_CONNECTIONS  names mode: test each connection afterwards (default true)
  DRY_RUN           true = report only, change nothing
  SECRETS_JSON      JSON of connection passwords (the IDMC_CONNECTION_PASSWORDS secret)
  MIGRATION_EXCEL   path of the workbook (default cicd/migration_objects.xlsx)
  RESULTS_EXCEL     path of the results workbook (default migration_results.xlsx)
"""

import datetime
import json
import os
import sys
import urllib.parse

import openpyxl
import requests
import yaml
from openpyxl.styles import Alignment, Font, PatternFill

import idmc_cicd as core
from idmc_cicd import DeployError, Org, env_flag, log

EXCEL_PATH = os.environ.get('MIGRATION_EXCEL') or os.path.join(core.ROOT, 'cicd', 'migration_objects.xlsx')
RESULTS_PATH = os.environ.get('RESULTS_EXCEL') or os.path.join(core.ROOT, 'migration_results.xlsx')

# column key -> header prefixes (case-insensitive) accepted in the workbook
COLUMNS = [
    ('name', ('name in democentral', 'name')),
    ('migrate', ('migrate',)),
    ('update', ('update if different', 'update')),
    ('test', ('test after migrate', 'test')),
    ('target', ('existing name in dfactory', 'name in dfactory')),
    ('secret', ('password secret key', 'password secret')),
]


def load_config():
    with open(core.CONFIG_PATH, encoding='utf-8') as f:
        cfg = yaml.safe_load(f) or {}
    for k in ('connections', 'schedules', 'publish', 'tests'):
        cfg.setdefault(k, {})
    cfg['connections']['settings'] = cfg['connections'].get('settings') or {}
    cfg['connection_map'] = cfg.get('connection_map') or {}
    return cfg


def names_from_env():
    raw = os.environ.get('NAMES', '')
    return [n.strip() for n in raw.replace('\n', ',').split(',') if n.strip()]


def yes(value, default=False):
    if value is None or str(value).strip() == '':
        return default
    return str(value).strip().lower() in ('y', 'yes', 'true', '1', 'x')


def read_sheet(wb, title):
    """Rows of one sheet as dicts (name, migrate, update, test, target, secret, row)."""
    ws = next((wb[s] for s in wb.sheetnames if s.strip().lower() == title), None)
    if ws is None:
        return []
    rows = list(ws.iter_rows(values_only=True))
    for hi, r in enumerate(rows):
        heads = [str(c or '').strip().lower() for c in r]
        if any(h.startswith('name') for h in heads):
            break
    else:
        return []
    idx = {}
    for key, prefixes in COLUMNS:
        for i, h in enumerate(heads):
            if i not in idx.values() and any(h.startswith(p) for p in prefixes):
                idx[key] = i
                break
    if 'name' not in idx:
        return []

    def cell(r, key):
        i = idx.get(key)
        return r[i] if i is not None and i < len(r) else None

    out = []
    for n, r in enumerate(rows[hi + 1:], start=hi + 2):
        name = str(cell(r, 'name') or '').strip()
        if not name or name.startswith('#'):
            continue
        out.append({'name': name, 'row': n,
                    'migrate': yes(cell(r, 'migrate')),
                    'update': yes(cell(r, 'update')),
                    'test': yes(cell(r, 'test'), True),
                    'target': str(cell(r, 'target') or '').strip(),
                    'secret': str(cell(r, 'secret') or '').strip()})
    return out


def list_names(org, resource):
    r = org.v2('GET', '/' + resource)
    if r.status_code != 200:
        raise DeployError('listing ' + resource + 's in ' + org.label + ' failed: ' + r.text[:300])
    body = r.json()
    return sorted({o.get('name') for o in (body if isinstance(body, list) else []) if o.get('name')})


def schedule_by_name(org, name):
    r = org.v2('GET', '/schedule/name/' + urllib.parse.quote(name, safe=''))
    if r.status_code == 200 and isinstance(r.json(), dict) and r.json().get('id'):
        return r.json()
    return None


def migrate_schedules(src, dst, names, update_existing, dry_run):
    """update_existing: True/False for all, or {name: True/False}."""
    log('\n== Schedules')
    rows = []
    for name in names:
        upd = update_existing.get(name, False) if isinstance(update_existing, dict) else update_existing
        s = schedule_by_name(src, name)
        row = {'name': name, 'target_hash': '-',
               'demo_hash': core.definition_hash(s, core.VOLATILE_SCHEDULE_KEYS) if s else '-'}
        if not s:
            row['action'] = 'not found in Demo'
        else:
            body = {k: v for k, v in s.items() if k not in core.VOLATILE_SCHEDULE_KEYS}
            t = schedule_by_name(dst, name)
            if not t:
                if dry_run:
                    row['action'] = 'missing (dry run: would create)'
                else:
                    c = dst.v2('POST', '/schedule', json=body)
                    row['action'] = 'created' if c.status_code == 200 else 'create failed: ' + c.text[:200]
            else:
                row['target_hash'] = core.definition_hash(t, core.VOLATILE_SCHEDULE_KEYS)
                if row['target_hash'] == row['demo_hash']:
                    row['action'] = 'exists, in sync'
                elif not upd:
                    row['action'] = 'exists, differs (not updated)'
                elif dry_run:
                    row['action'] = 'exists, differs (dry run: would update)'
                else:
                    u = dst.v2('POST', '/schedule/' + t['id'], json=body)
                    row['action'] = 'updated' if u.status_code == 200 else 'update failed: ' + u.text[:200]
        rows.append(row)
        log('  ' + name + ': ' + row['action'] + ' [demo ' + row['demo_hash'] + ' / dfactory ' + row['target_hash'] + ']')
    return rows


def migrate_connections(src, dst, cfg, secrets, dry_run, names, always_test, per_row=None):
    """per_row = {name: workbook row} to apply that row's update / test / target / secret settings."""
    settings, cmap = cfg['connections']['settings'], cfg['connection_map']
    for name, r in (per_row or {}).items():
        s = dict(settings.get(name) or {})
        s['update_existing'] = r['update']
        if r['test']:
            s['test_always'] = True
        else:
            s['test'] = False
        if r['secret']:
            s['password_secret'] = r['secret']
        settings[name] = s
        if r['target']:
            cmap[name] = r['target']
    report = {'connections': []}
    core.sync_connections(src, dst, [], cfg, dry_run, secrets, report, {}, None,
                          wanted={n: None for n in names}, always_test=always_test)
    return report['connections']


def write_summary(results, dry_run, orgs, source, error):
    md = '# IDMC migration: Demo → Dfactory ' + ('❌' if error else '✅') + '\n\n'
    md += '**Mode:** ' + ('dry run' if dry_run else 'migrate') + '  \n**Source:** ' + source
    if orgs:
        md += '  \n**Orgs:** ' + orgs
    md += '\n\n'
    if error:
        md += '> **Failed:** ' + error + '\n\n'
    for title, rows in results.items():
        md += '## ' + title + '\n' + core.table(rows, [('Name', 'name'), ('Result', 'action'), ('Demo hash', 'demo_hash'),
                                                       ('Dfactory hash', 'target_hash')]) + '\n'
    if not results:
        md += '_Nothing to migrate._\n'
    path = os.environ.get('GITHUB_STEP_SUMMARY')
    if path:
        with open(path, 'a', encoding='utf-8') as f:
            f.write(md)


def write_results_workbook(results, dry_run, orgs, source, error):
    font, bold = Font(name='Arial', size=10), Font(name='Arial', size=10, bold=True)
    head_fill = PatternFill('solid', start_color='D9E1F2')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Run'
    server, repo, run_id = (os.environ.get(k, '') for k in ('GITHUB_SERVER_URL', 'GITHUB_REPOSITORY', 'GITHUB_RUN_ID'))
    info = [('Run at (UTC)', datetime.datetime.utcnow().strftime('%Y-%m-%d %H:%M')),
            ('Mode', 'dry run' if dry_run else 'migrate'), ('Source', source), ('Orgs', orgs or '-'),
            ('Outcome', 'failed: ' + error if error else 'succeeded'),
            ('Commit', os.environ.get('GITHUB_SHA', '')[:7] or '-'),
            ('Run', server + '/' + repo + '/actions/runs/' + run_id if run_id else '-')]
    for r, (k, v) in enumerate(info, start=1):
        ws.cell(r, 1, k).font = bold
        ws.cell(r, 2, v).font = font
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 90
    for title, rows in results.items():
        sh = wb.create_sheet(title)
        heads = ['Name', 'Result', 'Demo hash', 'DFactory hash']
        for c, h in enumerate(heads, start=1):
            cell = sh.cell(1, c, h)
            cell.font, cell.fill = bold, head_fill
        for r, row in enumerate(rows, start=2):
            for c, key in enumerate(('name', 'action', 'demo_hash', 'target_hash'), start=1):
                sh.cell(r, c, row.get(key, '')).font = font
        for col, w in zip('ABCD', (42, 70, 16, 16)):
            sh.column_dimensions[col].width = w
        sh.freeze_panes = 'A2'
        for row in sh.iter_rows(min_row=2, max_col=2):
            row[1].alignment = Alignment(wrap_text=True, vertical='top')
    wb.save(RESULTS_PATH)
    log('Results written to ' + os.path.basename(RESULTS_PATH))


def main():
    cfg = load_config()
    names = names_from_env()
    source = (os.environ.get('SOURCE', '').strip().lower() or ('names' if names else 'excel'))
    otype = (os.environ.get('OBJECT_TYPE', '').strip().lower() or ('both' if source == 'excel' else ''))
    dry_run = env_flag('DRY_RUN', False)
    update_existing = env_flag('UPDATE_EXISTING', False)
    test = env_flag('TEST_CONNECTIONS', True)
    try:
        secrets = json.loads(os.environ.get('SECRETS_JSON') or '{}')
    except ValueError:
        secrets = {}

    logins, results, orgs, error = [], {}, '', None
    try:
        if source not in ('excel', 'names'):
            raise DeployError('SOURCE must be excel or names, not "' + source + '"')
        if otype not in ('connection', 'schedule', 'both'):
            raise DeployError('OBJECT_TYPE must be connection, schedule or both, not "' + otype + '"')
        conn_rows, sched_rows = {}, {}
        if source == 'excel':
            if not os.path.exists(EXCEL_PATH):
                raise DeployError(os.path.relpath(EXCEL_PATH, core.ROOT) + ' not found')
            wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
            if otype in ('connection', 'both'):
                conn_rows = {r['name']: r for r in read_sheet(wb, 'connections') if r['migrate']}
            if otype in ('schedule', 'both'):
                sched_rows = {r['name']: r for r in read_sheet(wb, 'schedules') if r['migrate']}
            log('Workbook ' + os.path.relpath(EXCEL_PATH, core.ROOT) + ': ' + str(len(conn_rows)) +
                ' connection(s) and ' + str(len(sched_rows)) + ' schedule(s) marked Migrate = Y')
            if not conn_rows and not sched_rows:
                log('Nothing to migrate')
                return
        else:
            if otype == 'both':
                raise DeployError('with typed names, choose connection or schedule')
            if not names:
                raise DeployError('no names given')
        log('Migrating' + (' (dry run)' if dry_run else ''))

        login_url = os.environ['IICS_LOGIN_URL']
        src = Org('Demo', login_url, os.environ['IICS_USERNAME'], os.environ['IICS_PASSWORD'])
        logins.append(src)
        dst = Org('Dfactory', login_url, os.environ['UAT_IICS_USERNAME'], os.environ['UAT_IICS_PASSWORD'])
        logins.append(dst)
        orgs = src.org_name + ' → ' + dst.org_name
        core.check_orgs(src, dst, cfg)

        if source == 'excel':
            if conn_rows:
                results['Connections'] = migrate_connections(src, dst, cfg, secrets, dry_run, list(conn_rows),
                                                             False, conn_rows)
            if sched_rows:
                results['Schedules'] = migrate_schedules(src, dst, list(sched_rows),
                                                         {n: r['update'] for n, r in sched_rows.items()}, dry_run)
        else:
            if names == ['*']:
                names = list_names(src, otype)
                log('All ' + otype + 's in Demo: ' + str(len(names)))
            if otype == 'connection':
                cfg['connections']['update_existing'] = update_existing
                results['Connections'] = migrate_connections(src, dst, cfg, secrets, dry_run, names, test)
            else:
                results['Schedules'] = migrate_schedules(src, dst, names, update_existing, dry_run)

        bad = [r for rows in results.values() for r in rows
               if r['action'].startswith('not found') or 'failed:' in r['action'].replace('test failed:', '')
               or 'mapped name' in r['action']]
        if bad:
            error = 'could not migrate: ' + ', '.join(r['name'] + ' (' + r['action'] + ')' for r in bad)
        log('\n' + ('Migration failed: ' + error if error else 'Migration ' + ('dry run ' if dry_run else '') + 'completed'))
    except (DeployError, KeyError, requests.RequestException) as e:
        error = str(e)
        log('\nMigration failed: ' + error)
    finally:
        write_summary(results, dry_run, orgs, source, error)
        write_results_workbook(results, dry_run, orgs, source, error)
        for o in logins:
            o.logout()
    sys.exit(0 if error is None else 1)


if __name__ == '__main__':
    main()

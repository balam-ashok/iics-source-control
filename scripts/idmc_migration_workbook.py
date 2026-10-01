"""
Build or refresh cicd/migration_objects.xlsx: every connection and schedule in the Demo org, with its status
in the Dfactory org (missing / exists, in sync / exists, differs). Read-only against both orgs.

Choices already made in the workbook (Migrate, Update if different, Test after migrate, Existing name in
DFactory, Password secret key, Notes) are kept, matched by name. Run by
.github/workflows/idmc_migration_workbook.yml, which commits the result.

Environment: IICS_LOGIN_URL, IICS_USERNAME / IICS_PASSWORD (Demo), UAT_IICS_USERNAME / UAT_IICS_PASSWORD
(Dfactory), MIGRATION_EXCEL (default cicd/migration_objects.xlsx).
"""

import datetime
import os
import sys

import openpyxl
import requests
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import idmc_cicd as core
import idmc_migrate_objects as mig
from idmc_cicd import DeployError, Org, log

HEAD_FILL = PatternFill('solid', start_color='D9E1F2')
INPUT_FILL = PatternFill('solid', start_color='FFF2CC')
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def font(**k):
    return Font(name='Arial', size=k.pop('size', 10), **k)


def by_name(org, resource):
    r = org.v2('GET', '/' + resource)
    if r.status_code != 200:
        raise DeployError('listing ' + resource + 's in ' + org.label + ' failed: ' + r.text[:300])
    body = r.json()
    return {o['name']: o for o in (body if isinstance(body, list) else []) if o.get('name')}


def statuses(src, dst, resource, volatile):
    demo, dfac = by_name(src, resource), by_name(dst, resource)
    out = {}
    for name, s in demo.items():
        t = dfac.get(name)
        if not t:
            out[name] = 'missing'
        elif core.definition_hash(s, volatile) == core.definition_hash(t, volatile):
            out[name] = 'exists, in sync'
        else:
            out[name] = 'exists, differs'
    return out


def previous_choices(path):
    """{sheet: {name: row dict}} from an existing workbook, so a refresh keeps what people entered."""
    if not os.path.exists(path):
        return {}
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for title in ('connections', 'schedules'):
        rows = {}
        ws = next((wb[s] for s in wb.sheetnames if s.strip().lower() == title), None)
        if ws is not None:
            notes_col = None
            head = [str(c.value or '').strip().lower() for c in ws[1]]
            if 'notes' in head:
                notes_col = head.index('notes')
            for r in mig.read_sheet(wb, title):
                row = ws[r['row']]
                r['notes'] = row[notes_col].value if notes_col is not None and notes_col < len(row) else None
                rows[r['name']] = r
        out[title] = rows
    return out


def read_me(ws, as_of):
    ws['A1'] = 'IDMC migration list: connections and schedules'
    ws['A1'].font = font(size=14, bold=True)
    ws['A2'] = ('DemoCentral to DFactory Inc DEV. Lists every connection and schedule in DemoCentral; status as of '
                + as_of + '. Refresh it with the "IDMC - Refresh migration workbook" workflow.')
    ws['A2'].font = font(color='595959')
    steps = [
        'How it works',
        '1. On the Connections and Schedules sheets, set Migrate = Y on the rows to move. Yellow cells are the ones to edit.',
        '2. Commit this file to main in iics-source-control as cicd/migration_objects.xlsx. The workflow '
        '"IDMC - Migrate connections or schedules" starts on its own.',
        '3. Open that run in GitHub Actions: its summary lists the result of every row, and migration_results.xlsx '
        'is attached to the run.',
        '4. To check first without changing anything: Actions > IDMC - Migrate connections or schedules > '
        'Run workflow, Source = excel, tick Dry run.',
        '5. After a successful run, set the migrated rows back to N, so the next commit does not repeat them '
        '(repeating is harmless: existing ones are only compared).',
    ]
    r = 4
    for i, s in enumerate(steps):
        ws.cell(r, 1, s).font = font(bold=(i == 0), size=(11 if i == 0 else 10))
        r += 1
    r += 1
    ws.cell(r, 1, 'Columns').font = font(bold=True, size=11)
    r += 1
    guide = [
        ('Column', 'Values', 'Meaning'),
        ('Name in DemoCentral', 'text', 'Exact name in DemoCentral. Add rows for new ones; rows whose name starts with # are ignored.'),
        ('Migrate', 'Y / N (blank = N)', 'Y = migrate this row on the next run.'),
        ('Update if different', 'Y / N (blank = N)', 'Y = overwrite the DFactory copy when its definition differs. '
                                                     'N = only create missing ones and report differences.'),
        ('Test after migrate', 'Y / N (blank = Y)', 'Connections only: run a connection test in DFactory afterwards.'),
        ('Existing name in DFactory', 'text, optional', 'Connections only: fill in when DFactory already has this '
                                                        'connection under a different name.'),
        ('Password secret key', 'text, optional', 'Connections only: key in the IDMC_CONNECTION_PASSWORDS GitHub '
                                                  'secret that holds the DFactory password.'),
        ('Status on ' + as_of, 'read only', 'Status at the last refresh: missing in DFactory, exists in sync, '
                                            'exists but differs, or not found in DemoCentral.'),
        ('Notes', 'text', 'Free text; the workflow ignores it.'),
    ]
    for i, row in enumerate(guide):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(r, c, v)
            cell.font = font(bold=(i == 0))
            cell.border = BOX
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if i == 0:
                cell.fill = HEAD_FILL
        r += 1
    r += 1
    ws.cell(r, 1, 'Example row (Connections sheet)').font = font(bold=True, size=11)
    r += 1
    example = [('Column', 'Example value'), ('Name in DemoCentral', 'Conn-BM-Oracle-Pharma'), ('Migrate', 'Y'),
               ('Update if different', 'N'), ('Test after migrate', 'Y'), ('Existing name in DFactory', '(blank)'),
               ('Password secret key', 'ORACLE_PHARMA_PWD'), ('Status on ' + as_of, 'missing'),
               ('Notes', 'Needed for the pharma demo')]
    for i, (k, v) in enumerate(example):
        for c, val in ((1, k), (2, v)):
            cell = ws.cell(r, c, val)
            cell.font = font(bold=(i == 0))
            cell.border = BOX
            if i == 0:
                cell.fill = HEAD_FILL
        r += 1
    r += 1
    ws.cell(r, 1, 'Good to know').font = font(bold=True, size=11)
    for n in ('Connections arrive without passwords unless a Password secret key is given; their test then fails '
              'until the password is set.',
              'Connection tests need an active Secure Agent in DFactory in the runtime environment the connection uses.',
              'A schedule migrated on its own is not attached to any mapping task; tasks are linked when they are '
              'deployed through the check-in flow.',
              'Names are matched exactly as they appear in DemoCentral, including spaces and case.'):
        r += 1
        ws.cell(r, 1, '- ' + n).font = font()
    for col, w in (('A', 30), ('B', 28), ('C', 90)):
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False


def data_sheet(wb, title, headers, rows, input_cols, yn_cols, status_col, widths):
    sh = wb.create_sheet(title)
    for c, h in enumerate(headers, start=1):
        cell = sh.cell(1, c, h)
        cell.font, cell.fill, cell.border = font(bold=True), HEAD_FILL, BOX
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for i, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            cell = sh.cell(i, c, v)
            cell.font = font(color='808080', italic=True) if c == status_col else font()
            cell.border = BOX
            if c in input_cols:
                cell.fill = INPUT_FILL
                if c in yn_cols:
                    cell.alignment = Alignment(horizontal='center')
    last = len(rows) + 1
    end = max(last, 500)
    dv = DataValidation(type='list', formula1='"Y,N"', allow_blank=True)
    dv.error, dv.errorTitle = 'Use Y or N', 'Y or N'
    sh.add_data_validation(dv)
    for c in yn_cols:
        dv.add(get_column_letter(c) + '2:' + get_column_letter(c) + str(end))
    sh.conditional_formatting.add('B2:B' + str(end),
                                  CellIsRule(operator='equal', formula=['"Y"'], fill=PatternFill('solid', start_color='C6EFCE'),
                                             font=Font(name='Arial', bold=True, color='006100')))
    st = get_column_letter(status_col)
    for text, color in (('missing', 'FCE4D6'), ('differs', 'FFEB9C'), ('not found', 'F2F2F2')):
        sh.conditional_formatting.add(st + '2:' + st + str(end),
                                      FormulaRule(formula=['ISNUMBER(SEARCH("' + text + '",' + st + '2))'],
                                                  fill=PatternFill('solid', start_color=color)))
    for c, w in enumerate(widths, start=1):
        sh.column_dimensions[get_column_letter(c)].width = w
    sh.row_dimensions[1].height = 30
    sh.freeze_panes = 'B2'
    sh.auto_filter.ref = 'A1:' + get_column_letter(len(headers)) + str(last)


ORDER = {'missing': 0, 'exists, differs': 1, 'exists, in sync': 2}


def build(path, conn_status, sched_status, prev, as_of):
    def merged(status, old):
        names = dict(status)
        for n in old:
            names.setdefault(n, 'not found in DemoCentral')
        return sorted(names.items(), key=lambda x: (ORDER.get(x[1], 3), x[0].lower()))

    yn = lambda v: 'Y' if v else 'N'
    pc, ps = prev.get('connections', {}), prev.get('schedules', {})
    wb = openpyxl.Workbook()
    read_me(wb.active, as_of)
    wb.active.title = 'Read me'
    conn_rows = []
    for name, st in merged(conn_status, pc):
        o = pc.get(name)
        conn_rows.append([name, yn(o and o['migrate']), yn(o and o['update']), yn(o['test'] if o else True),
                          (o and o['target']) or None, (o and o['secret']) or None, st, o and o.get('notes')])
    sched_rows = []
    for name, st in merged(sched_status, ps):
        o = ps.get(name)
        sched_rows.append([name, yn(o and o['migrate']), yn(o and o['update']), st, o and o.get('notes')])
    data_sheet(wb, 'Connections',
               ['Name in DemoCentral', 'Migrate', 'Update if different', 'Test after migrate', 'Existing name in DFactory',
                'Password secret key', 'Status on ' + as_of, 'Notes'],
               conn_rows, {2, 3, 4, 5, 6, 8}, {2, 3, 4}, 7, [46, 10, 12, 12, 26, 24, 21, 32])
    data_sheet(wb, 'Schedules',
               ['Name in DemoCentral', 'Migrate', 'Update if different', 'Status on ' + as_of, 'Notes'],
               sched_rows, {2, 3, 5}, {2, 3}, 4, [46, 10, 12, 21, 32])
    wb.active = 1
    wb.save(path)
    return len(conn_rows), len(sched_rows)


def main():
    cfg = mig.load_config()
    path = mig.EXCEL_PATH
    logins = []
    try:
        login_url = os.environ['IICS_LOGIN_URL']
        src = Org('Demo', login_url, os.environ['IICS_USERNAME'], os.environ['IICS_PASSWORD'])
        logins.append(src)
        dst = Org('Dfactory', login_url, os.environ['UAT_IICS_USERNAME'], os.environ['UAT_IICS_PASSWORD'])
        logins.append(dst)
        core.check_orgs(src, dst, cfg)
        conns = statuses(src, dst, 'connection', core.VOLATILE_CONNECTION_KEYS)
        scheds = statuses(src, dst, 'schedule', core.VOLATILE_SCHEDULE_KEYS)
        prev = previous_choices(path)
        kept = sum(1 for rows in prev.values() for r in rows.values() if r['migrate'])
        nc, ns = build(path, conns, scheds, prev, datetime.date.today().isoformat())
        log('Wrote ' + os.path.relpath(path, core.ROOT) + ': ' + str(nc) + ' connections, ' + str(ns) +
            ' schedules (' + str(kept) + ' rows still marked Migrate = Y)')
        for label, st in (('Connections', conns), ('Schedules', scheds)):
            counts = {}
            for v in st.values():
                counts[v] = counts.get(v, 0) + 1
            log('  ' + label + ': ' + ', '.join(k + ' ' + str(v) for k, v in sorted(counts.items())))
    except (DeployError, KeyError, requests.RequestException) as e:
        log('Refresh failed: ' + str(e))
        sys.exit(1)
    finally:
        for o in logins:
            o.logout()


if __name__ == '__main__':
    main()
